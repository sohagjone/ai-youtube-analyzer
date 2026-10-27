"""Editable Word (.docx) report built with python-docx."""
from __future__ import annotations

import io

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

from utils.charts import build_charts
from utils.data_processing import build_export_tables, prepare_df_for_export

MAX_ROWS = 30


def _cell(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "-"
    if isinstance(v, float):
        return f"{v:,.2f}" if abs(v) < 100 else f"{v:,.0f}"
    if isinstance(v, pd.Timestamp):
        return v.strftime("%Y-%m-%d")
    return str(v)[:200]


def _add_table(doc: Document, df: pd.DataFrame, cols: list[str] | None = None) -> None:
    df = prepare_df_for_export(df)
    if cols:
        df = df[[c for c in cols if c in df.columns]]
    df = df.head(MAX_ROWS)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns):
        table.rows[0].cells[i].text = str(c)
    for row in df.itertuples(index=False):
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = _cell(v)
    for r in table.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)


def build_docx(results: dict) -> bytes:
    doc = Document()
    meta, channel = results.get("meta", {}), results.get("channel", {})
    doc.add_heading("AI YouTube Channel Analyzer", 0)
    doc.add_paragraph(f"Channel: {channel.get('title', '')}")
    doc.add_paragraph(f"Analysis date (UTC): {str(meta.get('analysis_date', ''))[:19]} | Data mode: {meta.get('mode', '')}")
    if str(meta.get("mode", "")).startswith("SAMPLE"):
        doc.add_paragraph("SAMPLE DATA: every figure in this report is illustrative and not real YouTube data.").runs[0].bold = True
    tables, charts = build_export_tables(results), build_charts(results)

    def section(title, key, chart=None, cols=None, note=None):
        doc.add_heading(title, 1)
        if note:
            doc.add_paragraph(note)
        if chart and chart in charts:
            doc.add_picture(io.BytesIO(charts[chart]), width=Inches(6.2))
        if key in tables and len(tables[key]):
            _add_table(doc, tables[key], cols)
        elif not chart:
            doc.add_paragraph("No data available for this section.")

    rm = results.get("roadmap") or {}
    rm_note = (f"Targets are planning goals derived from peer benchmarks, not predictions. Health score: {rm.get('health_score')}/100. "
               f"Benchmark source: {rm.get('benchmark_source')}.") if rm.get("available") else None
    section("Growth roadmap", "Growth Roadmap", "Roadmap",
            ["Metric", "You", "Benchmark", "Status", "Day 30 target", "Day 60 target", "Day 90 target"], rm_note)
    section("Income estimate (estimate, not actual revenue)", "Income Estimate", None, ["Metric", "Value"],
            "Estimated from public data and editable assumptions; real income, watch hours and monetization status are private to the channel owner.")
    section("Income roadmap", "Income Roadmap", None, ["Step", "Milestone", "Target", "Today", "Estimated Date"])
    section("Time to monetization by scenario", "Income Timeline", None, None)
    section("Priority actions", "Roadmap Actions", None, ["Priority Rank", "Kpi", "Action", "How"])
    section("30-day action plan", "Daily Plan", None, ["Day", "Date", "Focus", "Task", "Deliverable"])
    section("1. Channel overview", "Channel Overview", note="Lifetime figures describe the whole channel; sample figures describe only the selected videos.")
    section("2. Video performance", "Video Performance", "Views timeline",
            ["Title", "Published (UTC)", "Views (public)", "Views per day", "Relative to median", "Format", "Performance flag"],
            "Public view counts only; low views are not automatically a failure.")
    section("3. Title analysis", "Title Patterns", None, None, "Patterns are observational, not causal.")
    section("4. Topic breakdown", "Topic Analysis", "Topics", None, f"Topic method: {meta.get('topic_method', '')}. Based on metadata, not full video content.")
    section("5. Competitor comparison", "Competitors", "Competitors",
            ["Role", "Channel", "Subscribers", "Median Views", "Uploads Per 30 Days", "Short Form Share", "Main Topics"])
    section("6. Content gap recommendations", "Content Gaps", "Content gaps",
            ["Rank", "Gap Type", "Topic", "Suggested Title", "Format", "Priority Score", "Confidence"],
            "A content gap is not proof of search demand. Suggested titles are recommendations, not guaranteed improvements.")
    section("7. Content calendar", "Content Calendar", None, ["Date", "Format", "Topic", "Working Title", "Basis"])
    doc.add_heading("8. Sources and limitations", 1)
    for s in results.get("sources", []):
        doc.add_paragraph(f"{s['type']}: {s['item']} - {s['detail']}", style="List Bullet")
    for lim in results.get("limitations", []):
        doc.add_paragraph(lim, style="List Number")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
