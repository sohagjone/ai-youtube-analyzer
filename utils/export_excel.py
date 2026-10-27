"""Excel workbook export (openpyxl): filters, number formats, frozen headers, clickable links."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.data_processing import build_export_tables, prepare_df_for_export

HEADER_FILL = PatternFill("solid", fgColor="D62828")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
MAX_CELL = 32000  # Excel limit is 32,767 characters


def _is_link_col(name: str) -> bool:
    n = name.lower()
    return "url" in n or n.endswith("link") or n == "detail"


def _number_format(name: str) -> str | None:
    n = name.lower()
    if any(k in n for k in ("rate", "share", "ratio")) and "per" not in n:
        return "0.0%"
    if "attainment" in n:
        return "0%"
    if "relative" in n or "score" in n or "similarity" in n or n in ("relevance", "demand evidence", "differentiation", "feasibility"):
        return "0.00"
    if any(k in n for k in ("views", "subscribers", "likes", "comments", "videos", "characters", "words")):
        return "#,##0"
    if "per day" in n or "per 30" in n or "per 1k" in n or "median" in n or "mean" in n:
        return "#,##0.0"
    return None


def _write_sheet(wb: Workbook, name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title=name[:31])
    df = prepare_df_for_export(df)
    cols = list(df.columns)
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, (col, val) in enumerate(zip(cols, row), start=1):
            if isinstance(val, float) and pd.isna(val):
                val = None
            elif isinstance(val, pd.Timestamp):
                val = val.to_pydatetime()
            elif isinstance(val, str):
                val = val[:MAX_CELL]
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if isinstance(val, str) and val.startswith(("http://", "https://")) and _is_link_col(col) and "\n" not in val:
                cell.hyperlink, cell.font = val, LINK_FONT
            elif isinstance(val, (int, float)) and not isinstance(val, bool):
                fmt = _number_format(col)
                if fmt:
                    cell.number_format = fmt
            elif hasattr(val, "year"):
                cell.number_format = "yyyy-mm-dd hh:mm"
            if isinstance(val, str) and len(val) > 60:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for c_idx, col in enumerate(cols, start=1):
        sample = [len(str(x)) for x in df[col].head(50).tolist() if x is not None]
        width = min(max([len(col)] + sample) + 2, 60)
        ws.column_dimensions[get_column_letter(c_idx)].width = max(width, 10)
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"


def build_workbook(results: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    tables = build_export_tables(results)
    for name in ("Channel Overview", "Growth Roadmap", "Roadmap Actions", "Daily Plan", "Video Performance", "Title Analysis", "Title Patterns", "Topic Analysis",
                 "Competitors", "Auto Competitors", "Content Gaps", "Content Calendar", "Sources and Notes"):
        if name in tables:
            _write_sheet(wb, name, tables[name])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
