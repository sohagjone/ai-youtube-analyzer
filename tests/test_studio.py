"""Tests for Trend Radar, video probe, editing advice, title/caption/hashtags, Shorts predictor, upload planner, copyright screen."""
import io
import json
import subprocess
import wave
from pathlib import Path
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import pytest

from tests.conftest import goto
from src import copy_generator as cg
from src import copyright_risk as cr
from src import editing_advice as ea
from src import shorts_predictor as sp
from src import studio, trends, upload_planner as up, video_probe as vp
from src.pipeline import run_analysis
from src.sample_data import SampleClient
from src.video_analysis import enrich_videos
from src.youtube_api import YouTubeAPIError, YouTubeClient
from utils.data_processing import videos_to_df

NOW = datetime(2026, 10, 27, tzinfo=timezone.utc)


# ------------------------------------------------------------------ helpers
def vids(n, views, days_ago_fn, dur=40, title_fn=None, channel="C"):
    rows = []
    for i in range(n):
        rows.append({"video_id": f"{channel}{i}", "title": (title_fn(i) if title_fn else f"Video {i} about espresso"), "url": f"https://youtu.be/{channel}{i}",
                     "published_at": (NOW - timedelta(days=days_ago_fn(i))).isoformat(), "duration_seconds": dur, "views": views(i), "likes": 10, "comments": 1,
                     "description": "", "tags": ["espresso"], "thumbnail_url": None, "category": None, "channel_id": channel, "channel_title": channel})
    return rows


class FakeClient:
    is_sample = False

    def __init__(self, bundles):
        self.bundles = bundles

    def fetch_channel_bundle(self, ref, n=30, date_from=None, date_to=None):
        if ref.value not in self.bundles:
            from src.youtube_api import ChannelNotFoundError
            raise ChannelNotFoundError("nope")
        return {"channel": {"title": ref.value, "channel_id": ref.value, "subscribers": 1000}, "videos": self.bundles[ref.value], "notes": []}


def fake_results(shorts_views=None, n_short=20, dur=30, topic_effect=False):
    """Minimal 'results' dict with Shorts history for the predictor."""
    rows = vids(n_short, (lambda i: shorts_views(i)) if shorts_views else (lambda i: 1000), lambda i: 20 + i * 3, dur=dur,
                title_fn=lambda i: f"{'Espresso' if i % 2 == 0 else 'Grinder'} tip number {i}")
    df = enrich_videos(videos_to_df(rows), NOW)
    df["topic"] = ["Espresso" if i % 2 == 0 else "Grinder" for i in range(len(df))][::-1] if False else df["title"].apply(lambda t: "Espresso" if t.startswith("Espresso") else "Grinder")
    topics = pd.DataFrame({"topic": ["Espresso", "Grinder"], "videos": [10, 10]})
    return {"videos": df, "topics": topics, "channel": {"subscribers": 1000}, "competitor_videos": {}, "competitor_channels": {}}


@pytest.fixture(scope="module")
def sample_results():
    return run_analysis(SampleClient(NOW), "@sample_own", [], sample_size=90, now=NOW, auto_competitors=10)


# ------------------------------------------------------------------ Trend Radar
def test_creator_radar_flags_pace_against_the_creators_own_baseline():
    # usual ~100 views/day; one recent video at ~600/day (6x) and one at ~100/day
    old = vids(12, lambda i: 100 * (20 + i * 5), lambda i: 20 + i * 5)
    recent = vids(1, lambda i: 600 * 5, lambda i: 5, title_fn=lambda i: "BREAKOUT video", channel="R") + vids(1, lambda i: 100 * 5, lambda i: 5, title_fn=lambda i: "Normal one", channel="Q")
    for k, v in enumerate(recent):
        v["video_id"] = f"new{k}"
    tbl, warn = trends.creator_radar(FakeClient({"@alice": old + recent}), ["@alice"], 30, 30, now=NOW)
    s = tbl.set_index("title")
    assert s.loc["BREAKOUT video", "status"] == "Breakout" and s.loc["BREAKOUT video", "pace_ratio"] > 5
    assert s.loc["Normal one", "status"] in ("Usual pace", "Above usual pace")
    assert tbl["pace_ratio"].is_monotonic_decreasing and warn == []
    assert list(trends.high_potential(tbl)["title"]) == ["BREAKOUT video"]


def test_creator_radar_handles_bad_input_missing_channels_and_limits():
    tbl, warn = trends.creator_radar(FakeClient({}), ["not a link", "@ghost"] + [f"@c{i}" for i in range(12)], now=NOW)
    assert tbl.empty and any("recognisable" in w for w in warn) and any("Only the first 10" in w for w in warn)


def test_trending_chart_and_terms_on_sample_client():
    chart = trends.trending_chart(SampleClient(NOW), "US", None, 20, NOW)
    assert 0 < len(chart) <= 20 and chart["views_per_hour"].is_monotonic_decreasing
    assert trends.learnings(chart.assign(status="x"))["available"] and len(trends.trend_terms(chart)) > 0
    assert "TikTok" in trends.PLATFORM_NOTE and "scraping" in trends.PLATFORM_NOTE.lower()


def test_live_client_most_popular_params():
    class Resp:
        status_code = 200

        def json(self):
            return {"items": []}

    class Sess:
        def __init__(self):
            self.params = None

        def get(self, url, params=None, timeout=None):
            self.params = params
            return Resp()

    c = YouTubeClient("AIzaFAKE", session=Sess())
    assert c.get_most_popular("bd", "20", 99) == []
    p = c.session.params
    assert p["chart"] == "mostPopular" and p["regionCode"] == "BD" and p["videoCategoryId"] == "20" and p["maxResults"] == 50


def test_trending_chart_error_is_an_api_error():
    class Bad:
        def get_most_popular(self, *a, **k):
            raise YouTubeAPIError("no chart")

    with pytest.raises(YouTubeAPIError):
        trends.trending_chart(Bad(), "US", "20")


# ------------------------------------------------------------------ video probe (real ffmpeg, synthetic file)
@pytest.fixture(scope="module")
def synthetic_video(tmp_path_factory):
    try:
        exe = vp.ffmpeg_exe()
    except vp.FFmpegMissing:
        pytest.skip("ffmpeg not available")
    d = tmp_path_factory.mktemp("vid")
    sr = 22050
    t = np.arange(sr * 12) / sr
    sig = np.zeros_like(t)
    for k in np.arange(0, 12, 0.5):  # 120 BPM clicks
        i, n = int(k * sr), int(0.05 * sr)
        sig[i:i + n] += np.sin(2 * np.pi * 880 * t[:n]) * np.exp(-np.linspace(0, 6, n))
    with wave.open(str(d / "a.wav"), "wb") as w:
        w.setnchannels(1), w.setsampwidth(2), w.setframerate(sr), w.writeframes((sig * 0.5 * 32767).astype(np.int16).tobytes())
    base = [exe, "-y", "-loglevel", "error"]
    colours = sum([["-f", "lavfi", "-i", f"color=c={c}:s=360x640:d=4:r=25"] for c in ("red", "blue", "green")], [])
    out = d / "t.mp4"
    subprocess.run(base + colours + ["-i", str(d / "a.wav"), "-filter_complex", "[0][1][2]concat=n=3:v=1:a=0[v]", "-map", "[v]", "-map", "3:a",
                                     "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", str(out)], check=True)
    silent = d / "silent.mp4"
    subprocess.run(base + ["-f", "lavfi", "-i", "color=c=red:s=640x360:d=3:r=25", "-c:v", "libx264", "-pix_fmt", "yuv420p", str(silent)], check=True)
    return {"mp4": out, "silent": silent, "dir": d}


def test_probe_measures_a_known_file(synthetic_video):
    r = vp.analyze_video(synthetic_video["mp4"])
    assert r["duration_seconds"] == pytest.approx(12, abs=0.2) and (r["width"], r["height"]) == (360, 640) and r["orientation"] == "vertical"
    assert r["cuts"]["count"] == 2 and r["cuts"]["avg_shot_seconds"] == pytest.approx(4.0, abs=0.1) and r["cuts"]["cuts_in_first_5s"] == 1
    assert r["tempo"]["bpm"] == pytest.approx(120, abs=6) and r["tempo"]["confidence"] > 0.3
    assert r["loudness"]["integrated_lufs"] < 0 and r["silence"]["silent_gaps"] == 0


def test_probe_handles_no_audio_and_unreadable_files(synthetic_video, tmp_path):
    r = vp.analyze_video(synthetic_video["silent"])
    assert not r["has_audio"] and r["loudness"] == {} and r["orientation"] == "horizontal"
    bad = tmp_path / "bad.mp4"
    bad.write_bytes(b"not a video at all")
    with pytest.raises(ValueError):
        vp.analyze_video(bad)


def test_analyze_bytes_always_deletes_its_temp_file(synthetic_video, monkeypatch):
    import os
    seen = []
    real = vp.analyze_video

    def spy(path):
        seen.append(path)
        return real(path)

    monkeypatch.setattr(vp, "analyze_video", spy)
    vp.analyze_bytes(synthetic_video["mp4"].read_bytes(), ".mp4")
    assert seen and not os.path.exists(seen[0])
    with pytest.raises(ValueError):
        vp.analyze_bytes(b"garbage", ".mp4")
    assert not os.path.exists(seen[-1])


# ------------------------------------------------------------------ editing advice
def probe_stub(**over):
    base = {"duration_seconds": 40.0, "width": 1920, "height": 1080, "orientation": "horizontal", "has_audio": True, "has_video": True,
            "loudness": {"integrated_lufs": -14.0, "true_peak_dbfs": -2.0}, "silence": {"silent_gaps": 0, "silent_seconds": 0, "starts_silent": False},
            "cuts": {"count": 10, "avg_shot_seconds": 3.5, "cuts_in_first_5s": 2}, "tempo": {"bpm": 100, "confidence": 0.5}}
    base.update(over)
    return base


CTX = {"available": True, "niche": "gaming", "preferred_format": "Short-form (<=3 min)", "best_duration_bucket": "<=1 min"}


def test_advice_flags_measured_problems_and_priorities():
    p = probe_stub(loudness={"integrated_lufs": -25.0, "true_peak_dbfs": 0.5}, silence={"silent_gaps": 4, "silent_seconds": 5, "starts_silent": True},
                   cuts={"count": 0, "avg_shot_seconds": 12.0, "cuts_in_first_5s": 0})
    a = ea.advise(p, CTX, "Short", "unlicensed")
    areas = list(a["area"])
    assert {"Framing", "Loudness", "Hook", "Pacing", "Music", "Music licence"} <= set(areas)
    assert a.iloc[0]["priority"] == "High" and a[a["priority"] == "High"].shape[0] >= 3
    assert any("-14" in x for x in a["advice"])
    assert (a["evidence"].str.len() > 0).all()


def test_advice_is_quiet_when_file_is_fine_and_works_without_a_file():
    ok = ea.advise(probe_stub(cuts={"count": 12, "avg_shot_seconds": 3.2, "cuts_in_first_5s": 2}), {**CTX, "preferred_format": "Long-form (>3 min)"}, "Long-form", "licensed")
    assert not (ok["priority"] == "High").any()
    none = ea.advise(None, CTX, "Short")
    assert "Upload a file" in set(none["area"])


def test_music_recommendation_uses_niche_and_compares_detected_tempo():
    rec = ea.music_recommendation({"niche": "education"}, False, {"bpm": 180, "confidence": 0.6})
    assert (rec["bpm_low"], rec["bpm_high"]) == (70, 100) and rec["fit"] == "outside" and "licen" in rec["licence"].lower()
    assert "fit" not in ea.music_recommendation({"niche": "education"}, False, {"bpm": 180, "confidence": 0.1})  # low-confidence tempo ignored
    assert ea.music_recommendation({"niche": "nonsense"}, True)["niche"] == "nonsense"  # unknown niche falls back to general guidance


# ------------------------------------------------------------------ title / caption / hashtags
CGX = {"keywords": ["espresso", "coffee"], "trending_terms": ["cold brew"], "pattern_perf": {"How-to / guide": (2.0, 10), "Question": (0.5, 8)},
       "best_length_bucket": "30-50 chars", "own_titles": ["How to pull a perfect espresso shot"], "top_tags": ["coffee", "espresso", "latte art"]}


def test_title_score_rewards_channel_patterns_and_penalises_duplicates_and_hype():
    good = cg.score_title("How to dial in espresso at home", CGX)
    weak = cg.score_title("Is espresso worth it?", CGX)
    dup = cg.score_title("How to pull a perfect espresso shot", CGX)
    hype = cg.score_title("SHOCKING espresso SECRET you won't believe!!", CGX)
    assert good["fit_score"] > weak["fit_score"] and dup["safety_pts"] == 0 and "similar" in dup["notes"] and hype["safety_pts"] < 15 and "wording" in hype["notes"]
    assert all(0 <= x["fit_score"] <= 100 for x in (good, weak, dup, hype))


def test_pattern_effect_is_shrunk_for_small_samples():
    big = cg.score_title("How to brew", {**CGX, "pattern_perf": {"How-to / guide": (3.0, 40)}})["pattern_pts"]
    small = cg.score_title("How to brew", {**CGX, "pattern_perf": {"How-to / guide": (3.0, 1)}})["pattern_pts"]
    assert big > small


def test_rank_titles_adds_user_draft_and_ai_titles_are_scored_too():
    df = cg.rank_titles("cold brew", CGX, "Simple ratios.", "Short", extra=["My own AI title about cold brew"])
    assert df["fit_score"].is_monotonic_decreasing and "AI" in set(df["source"]) and len(df) <= 8
    assert not df["title"].duplicated().any()


def test_hashtags_respect_limits_format_and_syntax():
    h = cg.hashtags("cold brew coffee", CGX, "Short", limit=20)
    assert len(h["hashtags"]) <= cg.MAX_HASHTAG_SAFE < cg.MAX_HASHTAGS and "#Shorts" in h["hashtags"] and h["above_title"] == h["hashtags"][:3]
    assert all(t.startswith("#") and " " not in t and t[1:].isalnum() for t in h["hashtags"])
    assert "#Shorts" not in cg.hashtags("cold brew coffee", CGX, "Long-form")["hashtags"] and len(h["hidden_tags"]) <= 480


def test_caption_uses_placeholders_instead_of_inventing_facts():
    c = cg.caption("cold brew", "Cold brew at home", "", "Long-form", ["#ColdBrew"])
    assert "[Write 2 sentences" in c and "[track] by [artist]" in c and "#ColdBrew" in c and "0:00 Intro" in c
    s = cg.caption("cold brew", "Cold brew at home", "Real summary.", "Short", ["#ColdBrew"], music_credit="x")
    assert "Real summary." in s and "Chapters" not in s


def test_ai_titles_failure_and_success_paths():
    class P:
        def complete_json(self, *a, **k):
            return ["A", " B ", 5, ""]

    class Empty:
        def complete_json(self, *a, **k):
            return []

    assert cg.ai_titles(P(), "t", "s", CGX, "Short") == ["A", "B"]
    with pytest.raises(ValueError):
        cg.ai_titles(Empty(), "t", "s", CGX, "Short")


# ------------------------------------------------------------------ Shorts predictor
def test_baseline_only_when_views_are_flat_and_range_collapses():
    r = fake_results(lambda i: 1000)
    o = sp.predict_short(r, "Espresso tip", 30, 2)
    assert o["available"] and o["estimate_mid"] == 1000 and o["estimate_low"] == 1000 and o["estimate_high"] == 1000 and not o["factors_used"]


def test_range_matches_the_spread_of_history_and_is_ordered():
    rng = np.random.default_rng(1)
    v = (rng.lognormal(7, 0.8, 30)).astype(int) + 1
    r = fake_results(lambda i: int(v[i]), n_short=30)
    o = sp.predict_short(r, "Espresso tip", 30)
    assert o["estimate_low"] <= o["estimate_mid"] <= o["estimate_high"] and o["confidence"] == "High" and o["backtest"]["n"] == 30
    assert 0.2 <= o["backtest"]["share_inside_range"] <= 0.8  # an honest IQR range should hold roughly half of past outcomes


def test_factors_are_used_only_when_they_beat_the_baseline_in_backtest():
    # topic strongly determines views -> factors help
    r = fake_results(lambda i: 5000 if i % 2 == 0 else 500, n_short=40)
    hi = sp.predict_short(r, "Espresso tip", 30)
    lo = sp.predict_short(r, "Grinder tip", 30)
    assert hi["factors_used"] and hi["backtest"]["factors_help"] and hi["estimate_mid"] > 2 * lo["estimate_mid"] and lo["estimate_mid"] < hi["baseline"] < hi["estimate_mid"]  # adjustments are clamped to 0.6-1.6x
    # nothing explains views -> baseline only, stated in the notes
    rng = np.random.default_rng(3)
    noise = rng.lognormal(7, 1.0, 40).astype(int)
    r2 = fake_results(lambda i: int(noise[i]), n_short=40)
    o2 = sp.predict_short(r2, "Espresso tip", 30)
    assert not o2["factors_used"] and any("did NOT beat" in n for n in o2["notes"])


def test_unavailable_with_too_few_shorts_and_never_negative():
    r = fake_results(lambda i: 100, n_short=2)
    o = sp.predict_short(r, "x", 30)
    assert not o["available"] and "at least 3" in o["reason"]
    r["videos"]["format"] = "Long-form (>3 min)"
    assert sp.predict_short(r, "x", 30)["n_shorts"] == 0


def test_peer_fallback_is_low_confidence_and_wide():
    r = fake_results(lambda i: 100, n_short=1)
    peers, chans = {}, {}
    for k in range(4):
        df = enrich_videos(videos_to_df(vids(10, lambda i: 2000, lambda i: 30 + i, channel=f"P{k}")), NOW)
        peers[f"P{k}"], chans[f"P{k}"] = df, {"subscribers": 10_000}
    r.update(competitor_videos=peers, competitor_channels=chans)
    o = sp.predict_short(r, "x", 30)
    assert o["available"] and o["confidence"] == "Low" and o["estimate_high"] / o["estimate_low"] >= 8
    assert o["estimate_mid"] == pytest.approx(0.2 * 1000)  # 2000 views / 10k subs = 0.2 per sub x 1,000 subs


# ------------------------------------------------------------------ upload planner
def timed_df(slots):
    """slots: list of (weekday 0-6, hour UTC, views, count)."""
    rows = []
    base = NOW - timedelta(days=200)
    k = 0
    for wd, hr, views, count in slots:
        for j in range(count):
            d = base + timedelta(days=(wd - base.weekday()) % 7 + 7 * j)
            d = d.replace(hour=hr, minute=0, second=0, microsecond=0)
            rows.append({"video_id": f"s{k}", "title": f"v{k}", "url": "", "published_at": d.isoformat(), "duration_seconds": 600, "views": views, "likes": 1, "comments": 1,
                         "description": "", "tags": [], "thumbnail_url": None, "category": None, "channel_id": "c", "channel_title": "c"})
            k += 1
    return enrich_videos(videos_to_df(rows), NOW)


def test_slot_table_shrinks_thin_slots_so_they_cannot_beat_well_evidenced_ones():
    df = timed_df([(1, 15, 1000, 8), (4, 9, 5000, 1), (2, 12, 1000, 6)])  # Tue 15h: 8 videos; Fri 09h: one lucky outlier
    t = up.slot_table(df, "UTC")
    tue = t[(t.weekday == 1) & (t.block == 15)].iloc[0]
    fri = t[(t.weekday == 4) & (t.block == 9)].iloc[0]
    assert tue.own_videos == 8 and fri.own_videos == 1 and fri.evidence == "none" and fri.score < fri.own_median_rel  # shrunk toward 1
    assert (t["score"] > 0).all() and len(t) == 56


def test_best_slots_need_evidence_and_avoid_list_needs_three_videos():
    df = timed_df([(1, 15, 3000, 6), (3, 9, 300, 4), (5, 21, 1000, 6)])
    rec = up.recommend(up.slot_table(df, "UTC"))
    assert rec["best"].iloc[0]["day"] == "Tue" and (rec["best"]["evidence"] != "none").all()
    assert list(rec["avoid"]["day"]) == ["Thu"]


def test_timezone_conversion_moves_slots():
    df = timed_df([(1, 20, 3000, 6), (2, 12, 1000, 6)])
    best_utc = up.recommend(up.slot_table(df, "UTC"))["best"].iloc[0]
    best_dhaka = up.recommend(up.slot_table(df, "Asia/Dhaka"))["best"].iloc[0]  # UTC+6: Tue 20:00 UTC -> Wed 02:00 local
    assert (best_utc["day"], best_utc["block"]) == ("Tue", 18) and (best_dhaka["day"], best_dhaka["block"]) == ("Wed", 0)
    assert up.tz_of("Not/AZone").key == "UTC"


def test_reliability_levels():
    assert up.reliability(timed_df([(1, 9, 100, 5)]))[0] == "Low" and up.reliability(timed_df([(0, 9, 100, 11), (1, 9, 100, 11), (2, 9, 100, 11)]))[0] == "Moderate"


def test_plan_respects_spacing_count_lead_time_and_idea_order():
    df = timed_df([(w, 15, 1000 + 200 * w, 4) for w in range(7)])
    slots = up.slot_table(df, "UTC")
    now = datetime(2026, 10, 28, 9, tzinfo=ZoneInfo("UTC"))
    plan = up.plan_uploads(slots, 8, [("Idea A", "Long-form"), ("Idea B", "Short")], now, horizon_days=14)
    times = pd.to_datetime(plan["publish_at_local"])
    assert len(plan) == round(8 / 30 * 7 * 14 / 7) and times.is_monotonic_increasing and times.min() >= pd.Timestamp("2026-10-28 21:00")
    assert (times.diff().dropna() >= pd.Timedelta(hours=20)).all() and list(plan["idea"][:2]) == ["Idea A", "Idea B"]
    assert up.plan_uploads(slots, 30, [], now, horizon_days=7).shape[0] <= 7


# ------------------------------------------------------------------ copyright / reuse risk
CLEAN = {"music": "own", "footage": "all_original", "commentary": False}


def test_clean_declaration_is_low_risk_but_never_a_clearance():
    r = cr.screen("My morning espresso routine", "How I make coffee", [], CLEAN, ["Another title"])
    assert r["level"] == "Low" and r["risk_score"] == 0 and "not a clearance" in r["summary"] and "Content ID" in r["disclaimer"] and "not legal advice" in r["disclaimer"]


def test_unlicensed_music_and_others_clips_are_high_and_commentary_reduces_clip_risk():
    worst = cr.screen("t", "d", [], {"music": "unlicensed", "footage": "others_clips", "commentary": False})
    mitigated = cr.screen("t", "d", [], {"music": "own", "footage": "others_clips", "commentary": True})
    plain = cr.screen("t", "d", [], {"music": "own", "footage": "others_clips", "commentary": False})
    assert worst["level"] == "High" and worst["findings"].iloc[0]["severity"] == "High" and worst["risk_score"] >= 60
    assert mitigated["risk_score"] < plain["risk_score"]


def test_text_signals_myth_disclaimers_and_duplicates():
    r = cr.screen("Full movie lyrics", "no copyright intended, fair use", [], CLEAN, ["Full movie lyrics!"], ["Full movie lyrics"])
    issues = " ".join(r["findings"]["issue"]).lower()
    assert "rights-sensitive" in issues and "disclaimer" in issues and "similar to your earlier" in issues and "another channel" in issues
    assert r["findings"]["severity"].tolist() == sorted(r["findings"]["severity"], key=["High", "Medium", "Low"].index)


def test_reuse_and_mass_production_flags_and_score_cap():
    r = cr.screen("t", "d", [], {**CLEAN, "reupload": True, "ai_generated": True, "templated": True, "lyrics_or_scripts": True, "music": "unlicensed", "footage": "others_clips"})
    assert r["risk_score"] == 100 and r["level"] == "High"
    t = cr.screen("t", "d", [], CLEAN, [], None, repeated_structure_share=0.8)
    assert "repetitive" in t["findings"].iloc[0]["issue"].lower()


# ------------------------------------------------------------------ orchestrator, exports, app
def test_run_studio_end_to_end_for_a_long_video_and_a_short(sample_results):
    base = {"title": "My cold brew routine", "topic": "cold brew coffee", "summary": "Simple.", "tz": "Asia/Dhaka", "music": "licensed", "footage": "all_original", "goal": "engagement"}
    long_ = studio.run_studio(sample_results, {**base, "format": "Long-form"}, None, None, ["espresso"], now_local=datetime(2026, 10, 28, 9, tzinfo=ZoneInfo("Asia/Dhaka")))
    assert long_["format"] == "Long-form" and not long_["shorts"]["available"] and long_["copyright"]["level"] == "Low"
    assert len(long_["titles"]) >= 2 and long_["titles"].iloc[0]["source"] == "Your draft" and long_["caption"].startswith(long_["titles"].sort_values("fit_score", ascending=False).iloc[0]["title"])
    short = studio.run_studio(sample_results, {**base, "format": "Short", "weekday": 2}, {"duration_seconds": 30.0, "orientation": "vertical", "has_audio": True, "has_video": True,
                                                                                         "width": 1080, "height": 1920, "cuts": {"avg_shot_seconds": 2.0, "cuts_in_first_5s": 2},
                                                                                         "loudness": {"integrated_lufs": -14.0}, "silence": {}, "tempo": {}}, None, None)
    assert short["format"] == "Short" and short["shorts"]["available"] and "#Shorts" in short["hashtags"]["hashtags"]
    assert len(short["planner"]["plan"]) >= 1 and len(short["planner"]["best"]) >= 1


def test_auto_format_detection_from_the_file(sample_results):
    p = {"duration_seconds": 25.0, "orientation": "vertical", "has_audio": False, "has_video": True, "width": 1080, "height": 1920, "cuts": {}, "loudness": {}, "silence": {}, "tempo": {}}
    assert studio.run_studio(sample_results, {"topic": "x y", "format": "Auto"}, p)["format"] == "Short"
    assert studio.run_studio(sample_results, {"topic": "x y", "format": "Auto"}, {**p, "duration_seconds": 600.0, "orientation": "horizontal"})["format"] == "Long-form"


def test_studio_exports_are_valid_and_contain_no_secrets(sample_results):
    import openpyxl

    from utils.data_processing import SECRET_PATTERNS, results_to_json_bytes
    from utils.export_studio import build_studio_workbook
    s = studio.run_studio(sample_results, {"title": "T", "topic": "cold brew", "format": "Short", "tz": "UTC"}, None, None)
    names = openpyxl.load_workbook(io.BytesIO(build_studio_workbook(s))).sheetnames
    assert {"Summary", "Editing Advice", "Titles", "Caption and Hashtags", "Copyright Risk", "Upload Plan"} <= set(names)
    blob = results_to_json_bytes(s)
    assert json.loads(blob)["format"] == "Short" and not any(p.search(blob.decode()) for p in SECRET_PATTERNS)


def test_app_trend_radar_and_video_studio_flows():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=180).run()
    [t for t in at.sidebar.text_input if "Your channel" in t.label][0].set_value("@sample_own")
    at.sidebar.button[0].click()
    at.run()
    names = list(at.radio(key="active_tab").options)
    assert "🔥 Trend Radar" in names and "🎬 Video Studio" in names and "Search" not in names and not at.exception
    goto(at, "Trend Radar")
    [b for b in at.button if b.label == "Show trending chart"][0].click()
    at.run()
    assert len(at.session_state["trend_chart"]) > 0 and not at.exception
    [t for t in at.text_area if t.label.startswith("Creators")][0].set_value("@rival_01\n@rival_02")
    [b for b in at.button if b.label == "Find high-potential videos"][0].click()
    at.run()
    assert at.session_state["trend_terms"] and not at.exception
    goto(at, "Video Studio")
    [t for t in at.text_input if t.label.startswith("Topic")][0].set_value("cold brew coffee")
    [b for b in at.button if b.label == "🎬 Analyse my video"][0].click()
    at.run()
    S = at.session_state["studio"]
    assert S["format"] == "Long-form" and len(S["titles"]) > 0 and not at.exception
    assert at.radio(key="active_tab").value.endswith("Video Studio")  # stays on the section after the rerun (the old tabs jumped back to Overview)
    assert any(w.value.startswith("Risk screen only") for w in at.warning)  # the disclaimer is always shown
    # empty submission is rejected, and a new analysis clears the studio
    [b for b in at.button if b.label == "Start a new analysis"][0].click()
    at.run()
    assert "studio" not in at.session_state or not at.session_state["studio"]
